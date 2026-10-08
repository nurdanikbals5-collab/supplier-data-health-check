"""M5 step 1: build the AI use-case prioritization workbook (scores left empty).

The scores, the weighted-score formula and the chart are added by hand in Excel,
so this script only writes the use-case list, the scoring guide and the weights.

Run:  ./.venv/bin/python src/m5_ai_use_cases_template.py
Out:  excel/AI_Use_Case_Prioritization.xlsx  (refuses to overwrite an existing file)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parents[1] / "excel" / "AI_Use_Case_Prioritization.xlsx"

# id, name, problem it solves, AI type, data it needs, data available in this project, status
USE_CASES = [
    ("UC1", "Duplicate supplier detection",
     "The same supplier exists in both ERPs under different names, so spend and quality data are split.",
     "Machine learning (fuzzy matching)", "Supplier master data (name, VAT, IBAN, city)", "Yes (M1)", "Idea"),
    ("UC2", "Complaint classification",
     "Quality complaints are free text; someone must read each one to find the category.",
     "Generative AI (LLM)", "Complaint texts + a few hand-labelled examples", "Yes (M3)", "Piloted (M3)"),
    ("UC3", "Master-data validation assistant",
     "Missing or invalid VAT numbers, IBANs and country codes block the S/4HANA migration.",
     "Rules + Generative AI (suggests fixes)", "Supplier master data + data-quality rules", "Yes (M1)", "Idea"),
    ("UC4", "Ticket triage assistant",
     "New data-quality tickets must be read, prioritised and assigned by hand.",
     "Generative AI (LLM)", "Ticket descriptions + past priority and owner", "Partly (ticket tracker)", "Idea"),
    ("UC5", "Late-delivery prediction",
     "Buyers learn about late deliveries only when the goods do not arrive.",
     "Machine learning (prediction)", "Purchase-order history with promised and actual dates", "Yes (PO lines)", "Idea"),
    ("UC6", "Supplier scorecard summary",
     "Buyers have no time to read KPI tables before a supplier meeting.",
     "Generative AI (LLM)", "Supplier KPIs (on-time delivery, PPM, spend)", "Yes (M2)", "Idea"),
    ("UC7", "S/4HANA help chatbot for users",
     "After go-live, users ask the same how-to questions about the new system again and again.",
     "Generative AI (chatbot on internal guides)", "Training material and process guides", "No", "Idea"),
    ("UC8", "Auto-correcting master-data agent",
     "Fixing master-data errors by hand is slow.",
     "AI agent (changes data in SAP by itself)", "Supplier master data + write access to SAP", "Partly", "Idea"),
]

CRITERIA = [
    # criterion, weight, higher score means, 1 = ..., 3 = ..., 5 = ...
    ("Business value", 0.40, "better",
     "Saves little time or money; few users", "Clear benefit for one team", "Large saving or risk reduction for many teams"),
    ("Data readiness", 0.20, "better",
     "Data does not exist or is very poor", "Data exists but needs cleaning", "Clean data is already available"),
    ("Effort", 0.20, "WORSE (more work)",
     "A few days, existing tools", "Some weeks, needs IT support", "Months, new system or integration"),
    ("Risk", 0.20, "WORSE (more risk)",
     "A wrong answer is harmless and easy to spot", "A wrong answer costs some time to fix", "A wrong answer changes real data or money without a check"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP


def main():
    if OUT.exists():
        raise SystemExit(f"{OUT.name} already exists - not overwriting your work.")
    wb = Workbook()

    # Sheet 1: use cases + empty score columns
    ws = wb.active
    ws.title = "Use Cases"
    headers = ["ID", "Use case", "Problem it solves", "AI type", "Data needed",
               "Data in project?", "Status", "Business value (1-5)", "Data readiness (1-5)",
               "Effort (1-5)", "Risk (1-5)", "Human approval needed?"]
    ws.append(headers)
    for uc in USE_CASES:
        ws.append(list(uc) + [None, None, None, None, None])
    widths = [6, 28, 48, 28, 36, 18, 14, 12, 12, 12, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=1, max_row=len(USE_CASES) + 1):
        for cell in row:
            cell.alignment = WRAP
    style_header(ws, len(headers))
    ws.freeze_panes = "C2"
    last = len(USE_CASES) + 1
    tbl = Table(displayName="tblUseCases", ref=f"A1:L{last}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tbl)
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                           showErrorMessage=True, errorTitle="Score 1-5",
                           error="Please enter a whole number from 1 to 5.")
    ws.add_data_validation(score)
    score.add(f"H2:K{last}")
    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(yes_no)
    yes_no.add(f"L2:L{last}")

    # Sheet 2: scoring guide + weights
    g = wb.create_sheet("Scoring Guide")
    g.append(["Criterion", "Weight", "A higher score is ...", "1 =", "3 =", "5 ="])
    for c in CRITERIA:
        g.append(list(c))
    g.append(["Total", "=SUM(B2:B5)"])
    for i, w in enumerate([18, 9, 20, 40, 40, 44], start=1):
        g.column_dimensions[g.cell(row=1, column=i).column_letter].width = w
    for row in g.iter_rows(min_row=1, max_row=6):
        for cell in row:
            cell.alignment = WRAP
    for r in range(2, 7):
        g.cell(row=r, column=2).number_format = "0%"
    style_header(g, 6)
    g["A8"] = "Effort and Risk count against a use case: in the weighted score they are turned around with (6 - score)."
    g["A9"] = "Scores are my own judgement for a simulated company with synthetic data, not measured values."
    g["A8"].font = g["A9"].font = Font(italic=True)

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
