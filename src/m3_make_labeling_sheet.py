"""M3 (step 1) - Pick 30 quality complaints for HUMAN labelling before any AI sees them.

The sample has distinct texts, about 5 clear cases per category and 6 ambiguous cases, in random order.
The answer key is NOT written into the sheet. The human labels later serve as the benchmark for the AI.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 3
ROOT = Path(__file__).resolve().parents[1]
qn = pd.read_csv(ROOT / "data" / "raw" / "quality_notifications.csv", dtype=str)
truth = pd.read_csv(ROOT / "data" / "raw" / "_truth_quality_categories.csv", dtype=str)
pool = qn.merge(truth, on="notification_id").drop_duplicates("description")

ambiguous = pool[pool.ambiguous == "True"].sample(6, random_state=SEED)
clear = pool[pool.ambiguous == "False"].groupby("true_category").sample(5, random_state=SEED).sample(24, random_state=SEED)
sample = pd.concat([ambiguous, clear]).sample(frac=1, random_state=SEED)

CATEGORIES = ["Leakage", "Dimensional", "Surface", "Packaging/Labeling", "Documentation", "Unclear"]
wb = Workbook()
ws = wb.active
ws.title = "Label"
ws.append(["#", "Notification ID", "Complaint text", "My category", "Note (optional)"])
for i, r in enumerate(sample.itertuples(), start=1):
    ws.append([i, r.notification_id, r.description, None, None])
for c in ws[1]:
    c.font = Font(bold=True)
for col, w in zip("ABCDE", [4, 15, 62, 20, 40]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
    row[0].alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"D2:D{len(sample) + 1}")

guide = wb.create_sheet("Guide")
rows = [
    ("Category", "Choose it when ...", "Example"),
    ("Leakage", "fluid or gas escapes, or a leak / pressure test fails", "Coolant leakage at brazed joint"),
    ("Dimensional", "size, shape or position does not match the drawing; the part does not fit",
     "Hole position shifted 0.4 mm"),
    ("Surface", "the surface is damaged or dirty: scratches, dents, corrosion, porosity, coating, contamination",
     "Corrosion spots found on delivery"),
    ("Packaging/Labeling", "box, pallet, container or label is wrong, damaged or missing; parts are mixed",
     "Wrong label on pallet"),
    ("Documentation", "a document is missing or wrong: certificate, test report, delivery note, PPAP",
     "Material certificate missing"),
    ("Unclear", "you cannot decide; write why in the Note column", ""),
    ("", "", ""),
    ("Rule", "Label what is wrong with the delivery. If two categories fit, pick the main problem and write the "
             "other one in the Note column.", ""),
    ("Rule", "Work alone and do not look at any answer key: your labels are the benchmark for the AI.", ""),
    ("", "German terms: Lötnaht = brazing seam, Maß = dimension, Kratzer = scratch, Etikett = label, "
         "Lieferschein = delivery note", ""),
]
for r in rows:
    guide.append(r)
for c in guide[1]:
    c.font = Font(bold=True)
for col, w in zip("ABC", [20, 90, 34]):
    guide.column_dimensions[col].width = w

out = ROOT / "excel" / "M3_Complaint_Labeling.xlsx"
if out.exists():
    raise SystemExit(f"{out} already exists (it holds the human labels) - not overwritten")
wb.save(out)
print(f"{len(sample)} complaints ({sample.ambiguous.eq('True').sum()} ambiguous) written to {out}")
print(sample.true_category.value_counts().to_string())
