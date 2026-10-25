"""M8 (input) - clean data file for the Power BI dashboard.

Power BI (web) reads plain values best, so this script copies the tickets from the Excel tracker I built
(excel/BU2_Ticket_Tracker.xlsx) and recalculates the formula columns for a fixed snapshot date:
Age (days) = Closed On - Created On, or Snapshot Date - Created On for open tickets
SLA Status = "Breached" if Age > SLA days, else "OK"   (same rules as the Excel formulas)

Run:  ./.venv/bin/python src/m8_export_powerbi_data.py [YYYY-MM-DD]
Out:  excel/powerbi/Tickets_PowerBI.xlsx  (one Excel table "Tickets")
"""
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "excel" / "BU2_Ticket_Tracker.xlsx"
OUT = ROOT / "excel" / "powerbi" / "Tickets_PowerBI.xlsx"
SNAPSHOT = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()

KEEP = ["Ticket ID", "Created On", "System", "Supplier", "Field", "DQ Dimension", "Issue", "Priority", "Status",
        "Owner", "Closed On", "Segment", "SLA (days)"]


def as_date(v):
    return v.date() if isinstance(v, datetime) else v


def main():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # the tracker has a slicer that openpyxl cannot read; we only read values
        ws = load_workbook(SRC, data_only=True)["Tickets"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    idx = {h: head.index(h) for h in KEEP}

    wb = Workbook()
    out = wb.active
    out.title = "Tickets"
    out.append(KEEP + ["Age (days)", "SLA Status", "Open", "Snapshot Date"])
    n = 0
    for r in rows[1:]:
        if not r[idx["Ticket ID"]]:
            continue
        created, closed = as_date(r[idx["Created On"]]), as_date(r[idx["Closed On"]])
        age = ((closed or SNAPSHOT) - created).days
        sla = r[idx["SLA (days)"]]
        is_open = r[idx["Status"]] != "Closed"
        out.append([as_date(r[idx[h]]) if h in ("Created On", "Closed On") else r[idx[h]] for h in KEEP]
                   + [age, "Breached" if age > sla else "OK", "Yes" if is_open else "No", SNAPSHOT])
        n += 1
    for row in out.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, date):
                c.number_format = "yyyy-mm-dd"
    t = Table(displayName="Tickets", ref=f"A1:Q{n + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    out.add_table(t)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"{n} tickets, snapshot {SNAPSHOT} -> {OUT}")


if __name__ == "__main__":
    main()
