"""M4 (input) - Supplier lookup table for the Excel lessons (XLOOKUP).

One row per source record (ERP_A vendor number or ERP_B supplier ID) with its golden record from M1
and its segment from M2. In Excel, the tickets look up the segment by Record ID.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
golden = pd.read_csv(ROOT / "data" / "output" / "golden_suppliers.csv", dtype=str)
segments = pd.read_csv(ROOT / "data" / "output" / "supplier_segments.csv", dtype=str)
segment_of = dict(zip(segments.golden_id, segments.segment))

wb = Workbook()
ws = wb.active
ws.title = "Suppliers"
ws.append(["Record ID", "System", "Golden ID", "Golden Name", "Segment"])
rows = 0
for g in golden.itertuples():
    for source in g.source_records.split("; "):
        system, record_id = source.split(":")
        ws.append([record_id, system, g.golden_id, g.name, segment_of[g.golden_id]])
        rows += 1
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    row[0].number_format = "@"  # keep leading zeros of SAP vendor numbers
for col, w in zip("ABCDE", [12, 8, 10, 26, 30]):
    ws.column_dimensions[col].width = w

out = ROOT / "excel" / "supplier_lookup.xlsx"
wb.save(out)
print(f"{rows} source records written to {out}")
