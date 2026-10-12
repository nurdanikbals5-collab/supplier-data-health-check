"""M7 (input) - starting workbook for mapping the project to SAP S/4HANA concepts.

The source fields, example values and project steps are filled in. The SAP columns
(target object, table.field, BP role, transformation rule, SAP app/transaction) are left
empty and are filled in by hand in Excel, using the SAP_Glossary sheet as a reference.

This is a conceptual mapping: no live SAP system is used.

Run:  ./.venv/bin/python src/m7_export_sap_mapping_template.py
Out:  excel/SAP_S4HANA_Mapping.xlsx  (refuses to overwrite an existing file)
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
RAW, OUT = ROOT / "data" / "raw", ROOT / "data" / "output"
XLSX = ROOT / "excel" / "SAP_S4HANA_Mapping.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TODO_FILL = PatternFill("solid", fgColor="FFF2CC")  # yellow = to be filled in by hand

# legacy field -> field in the golden record (M1)
GOLDEN_FIELD = {
    "LIFNR": "source_records", "NAME1": "name", "LAND1": "country", "ORT01": "city", "PSTLZ": "postcode",
    "STCEG": "vat", "IBAN": "iban", "ZTERM": "terms", "ERDAT": "",
    "supplier_code": "source_records", "company_name": "name", "country": "country", "city": "city",
    "zip": "postcode", "vat_number": "vat", "bank_account": "iban", "payment_terms": "terms",
    "main_commodity": "",
}

PROCESS_STEPS = [
    ("M1", "Check supplier master data from two ERPs and merge duplicates into golden records"),
    ("M1", "Load the 120 golden supplier records into the new system"),
    ("M2", "Purchase order lines with promised date (basis for on-time delivery)"),
    ("M2", "Delivery date and received quantity (goods receipt)"),
    ("M2", "Rejected quantity and PPM per supplier"),
    ("M2", "Supplier KPIs and segments (on-time delivery, PPM, spend)"),
    ("M3", "Quality complaints about a supplier (quality notifications)"),
    ("M4", "Track data-quality issues as tickets with priority and SLA"),
]

GLOSSARY = [
    ("SAP S/4HANA", "SAP's current ERP suite and the successor of SAP ECC.",
     "Target system of the simulated migration"),
    ("Business Partner (BP)", "In S/4HANA, suppliers and customers are maintained as business partners "
     "(transaction BP). This replaces the separate vendor master transactions of SAP ECC (e.g. XK01).",
     "Each golden record becomes one business partner"),
    ("BP role FLVN00", "Supplier role for Financial Accounting: company code data such as payment terms.",
     "Payment terms (terms)"),
    ("BP role FLVN01", "Supplier role for Purchasing: purchasing organization data such as order currency "
     "and purchasing payment terms.", "Supplier used in purchase orders"),
    ("CVI (Customer/Vendor Integration)", "Keeps the business partner in sync with the classic supplier "
     "tables LFA1, LFB1 and LFM1.", "Why the ERP A field names still matter in S/4HANA"),
    ("LFA1", "Supplier general data. Fields: LIFNR (supplier number), NAME1 (name), LAND1 (country), "
     "ORT01 (city), PSTLZ (postal code), STCEG (VAT registration number), ERDAT (created on).",
     "ERP A already uses these field names"),
    ("LFB1", "Supplier data per company code, e.g. ZTERM (payment terms key).", "Payment terms"),
    ("LFM1", "Supplier data per purchasing organization, e.g. ZTERM (payment terms in purchasing).",
     "Payment terms"),
    ("LFBK / IBAN", "Supplier bank details; the IBAN is stored with the bank details "
     "(bank details tab of the business partner).", "iban / bank_account"),
    ("Migration Cockpit", "SAP S/4HANA tool for loading legacy data ('Migrate Your Data' app) with "
     "predefined migration objects, e.g. 'Supplier'.", "Loading the golden records"),
    ("Master Data Governance (MDG)", "SAP solution for central master data maintenance with approval "
     "workflows and data quality rules.", "Compare with the M1 rules and the M4 ticket process"),
    ("Purchase order (EKKO / EKPO)", "Purchase order header (EKKO) and items (EKPO); display with ME23N. "
     "Standard purchase order numbers often start with 45.", "po_number 45..."),
    ("Schedule line (EKET)", "Delivery schedule of a PO item; EINDT is the planned delivery date.",
     "promised_date"),
    ("Goods receipt (MIGO, movement type 101)", "Posting of received goods; creates a material document "
     "(table MATDOC in S/4HANA).", "delivery_date, qty_received"),
    ("QM notification (QM01 / QM02 / QM03)", "Quality notification to record a problem; create / change / "
     "display. Notification type Q2 = complaint against a supplier. Header table QMEL.",
     "quality_notifications.csv (QN...)"),
    ("Supplier evaluation", "S/4HANA can score suppliers with criteria such as price, quality and delivery "
     "time.", "Compare with the M2 supplier KPIs"),
]


def header(ws):
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def add_table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def mark_todo(ws, cols, nrows):
    for col in cols:
        for r in range(2, nrows + 2):
            ws[f"{col}{r}"].fill = TODO_FILL


def main():
    if XLSX.exists():
        raise SystemExit(f"{XLSX.name} already exists - not overwriting your work.")

    wb = Workbook()

    # 1) Field mapping: one row per legacy field
    ws = wb.active
    ws.title = "Field_Mapping"
    ws.append(["Source system", "Source field", "Example value", "Golden record field (M1)",
               "S/4HANA object", "SAP table.field", "BP role", "Transformation rule"])
    rows = 0
    for system, file in [("ERP_A", "erp_a_vendors.csv"), ("ERP_B", "erp_b_suppliers.csv")]:
        df = pd.read_csv(RAW / file, dtype=str)
        for col in df.columns:
            example = df[col].dropna().iloc[0] if df[col].notna().any() else ""
            ws.append([system, col, example, GOLDEN_FIELD.get(col, ""), "", "", "", ""])
            rows += 1
    header(ws)
    mark_todo(ws, "EFGH", rows)
    add_table(ws, "tblFieldMap", f"A1:H{rows + 1}")
    for col, w in zip("ABCDEFGH", [12, 16, 26, 22, 22, 18, 12, 40]):
        ws.column_dimensions[col].width = w

    # 2) Process mapping: one row per project step
    ws = wb.create_sheet("Process_Mapping")
    ws.append(["Module", "Step in this project", "SAP S/4HANA app / transaction", "SAP object or table",
               "Note"])
    for module, step in PROCESS_STEPS:
        ws.append([module, step, "", "", ""])
    header(ws)
    mark_todo(ws, "CDE", len(PROCESS_STEPS))
    add_table(ws, "tblProcessMap", f"A1:E{len(PROCESS_STEPS) + 1}")
    for col, w in zip("ABCDE", [9, 60, 30, 22, 40]):
        ws.column_dimensions[col].width = w

    # 3) Data-quality rules from M1 -> later used as transformation rules
    ws = wb.create_sheet("DQ_Rules")
    issues = pd.read_csv(OUT / "dq_issues.csv", dtype=str)
    # "Possible duplicate of 0000100192 (Same VAT ID)" -> one rule for all duplicates
    issues["rule"] = issues.rule.str.replace(r"Possible duplicate of \S+ \((.*)\)", r"Possible duplicate (\1)",
                                             regex=True)
    rules = (issues.groupby(["system", "field", "rule", "suggested_action"]).size()
             .reset_index(name="issues").sort_values(["system", "field"]))
    ws.append(["System", "Field", "Rule (M1)", "Suggested action (M1)", "Issues"])
    for r in rules.itertuples(index=False):
        ws.append(list(r))
    header(ws)
    add_table(ws, "tblDQRules", f"A1:E{len(rules) + 1}")
    for col, w in zip("ABCDE", [10, 16, 40, 40, 8]):
        ws.column_dimensions[col].width = w

    # 4) Reference
    ws = wb.create_sheet("SAP_Glossary")
    ws.append(["SAP term", "What it is", "Where it appears in this project"])
    for row in GLOSSARY:
        ws.append(list(row))
    header(ws)
    add_table(ws, "tblGlossary", f"A1:C{len(GLOSSARY) + 1}")
    for col, w in zip("ABC", [32, 80, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    XLSX.parent.mkdir(exist_ok=True)
    wb.save(XLSX)
    print(f"saved {XLSX.relative_to(ROOT)}: {rows} fields, {len(PROCESS_STEPS)} steps, "
          f"{len(rules)} DQ rules, {len(GLOSSARY)} glossary terms")


if __name__ == "__main__":
    main()
