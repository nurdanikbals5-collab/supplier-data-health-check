"""M4 (input) - Turn every data-quality issue from M1 into a ticket for the Excel ticket tracker.

Ticket workflow fields (created date, owner, status) are SIMULATED so that the Excel exercises
(ageing, SLA, KPIs) have realistic data. The issues themselves come from M1.
"""
import random
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(11)
ROOT = Path(__file__).resolve().parents[1]
dq = pd.read_csv(ROOT / "data" / "output" / "dq_issues.csv", dtype=str)

PRIORITY = {"High": "High", "Medium": "Medium", "Low": "Low"}
OWNERS = ["A. Weber", "L. Rossi", "M. Chen", "S. Kaya", "J. Novak"]

wb = Workbook()
ws = wb.active
ws.title = "Tickets"
headers = ["Ticket ID", "Created On", "System", "Record ID", "Supplier", "Field", "DQ Dimension", "Issue",
           "Suggested Action", "Priority", "Status", "Owner", "Closed On"]
ws.append(headers)
for r in dq.itertuples():
    created = date(2026, 9, 1) + timedelta(days=random.randint(0, 33))
    status = random.choices(["Open", "In Progress", "Waiting", "Closed"], [3, 3, 1, 4])[0]
    closed = min(created + timedelta(days=random.randint(1, 10)), date(2026, 10, 5)) if status == "Closed" else None
    ws.append([r.issue_id.replace("DQ-", "T-"), created, r.system, r.record_id, r.supplier_name, r.field, r.dimension,
               r.rule, r.suggested_action, PRIORITY[r.severity], status, random.choice(OWNERS), closed])
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    row[1].number_format = "yyyy-mm-dd"
    row[12].number_format = "yyyy-mm-dd"
for col, w in zip("ABCDEFGHIJKLM", [9, 12, 8, 12, 22, 10, 13, 42, 38, 9, 12, 11, 12]):
    ws.column_dimensions[col].width = w

about = wb.create_sheet("About")
about.append(["BU2 Ticket Tracker - Supplier master-data migration (practice file)"])
about["A1"].font = Font(bold=True)
about.append(["Issues come from the M1 data-quality checks on a SYNTHETIC dataset."])
about.append(["Created date, status, owner and closed date are simulated for the Excel exercises. Owners are fictional."])

out = ROOT / "excel" / "BU2_Ticket_Tracker_start.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print(f"{len(dq)} tickets written to {out}")
