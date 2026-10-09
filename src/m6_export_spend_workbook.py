"""M6 (input) - starting workbook for the spend and payment-terms analysis in Excel.

One row per golden supplier with facts only (spend, commodity, segment, KPIs, payment term).
All calculations (shares, ABC class, cash effects) are added by hand in Excel.

Run:  ./.venv/bin/python src/m6_export_spend_workbook.py
Out:  excel/Spend_Analysis.xlsx  (refuses to overwrite an existing file)
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
RAW, OUT = ROOT / "data" / "raw", ROOT / "data" / "output"
XLSX = ROOT / "excel" / "Spend_Analysis.xlsx"

# payment terms used in both ERPs; "2% 14 days, net 30" = 2 % cash discount if paid within 14 days
TERMS = [("NT30", "Net 30", 30, 0.0, 0), ("NT45", "45 days", 45, 0.0, 0),
         ("NT60", "60 days net", 60, 0.0, 0), ("SK14", "2% 14 days, net 30", 30, 0.02, 14)]
TARGET_DAYS = 60  # assumption: company standard payment term

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def header(ws):
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    if XLSX.exists():
        raise SystemExit(f"{XLSX.name} already exists - not overwriting your work.")

    golden = pd.read_csv(OUT / "golden_suppliers.csv", dtype=str)
    kpis = pd.read_csv(OUT / "supplier_kpis.csv")
    po = pd.read_csv(RAW / "purchase_order_lines.csv", dtype={"supplier_id": str})

    # same legacy -> golden mapping as M2
    legacy_to_golden = {rec.split(":", 1)[1]: g.golden_id
                        for g in golden.itertuples() for rec in g.source_records.split("; ")}
    po["golden_id"] = po.supplier_id.map(legacy_to_golden)
    commodity = po.groupby("golden_id").commodity.agg(lambda s: s.mode().iat[0])
    first, last = pd.to_datetime(po.delivery_date).min(), pd.to_datetime(po.delivery_date).max()
    period_days = (last - first).days + 1

    df = kpis.merge(golden[["golden_id", "terms"]], on="golden_id", how="left")
    df["commodity"] = df.golden_id.map(commodity)
    df["terms"] = df.terms.fillna("missing")

    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    ws.append(["Golden ID", "Supplier", "Country", "Commodity", "Segment (M2)", "Spend (EUR)",
               "OTD %", "PPM", "Payment term"])
    for r in df.itertuples():
        ws.append([r.golden_id, r.name, r.country, r.commodity, r.segment, round(r.spend_eur, 2),
                   r.otd_pct, r.ppm, r.terms])
    n = len(df) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        row[5].number_format = "#,##0"
        row[6].number_format = "0.0"
        row[7].number_format = "#,##0"
    for col, w in zip("ABCDEFGHI", [10, 26, 8, 22, 32, 14, 8, 8, 13]):
        ws.column_dimensions[col].width = w
    header(ws)
    ws.freeze_panes = "C2"
    t = Table(displayName="tblSpend", ref=f"A1:I{n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)

    tr = wb.create_sheet("Terms")
    tr.append(["Code", "Description", "Net days", "Cash discount %", "Discount days"])
    for code, text, days, disc, ddays in TERMS:
        tr.append([code, text, days, disc, ddays])
    for r in range(2, len(TERMS) + 2):
        tr.cell(row=r, column=4).number_format = "0%"
    for col, w in zip("ABCDE", [8, 20, 10, 15, 14]):
        tr.column_dimensions[col].width = w
    header(tr)

    st = wb.create_sheet("Settings")
    st.append(["Setting", "Value", "Note"])
    st.append(["Data period: first delivery", first.date(), "Spend in this workbook covers this period only"])
    st.append(["Data period: last delivery", last.date(), ""])
    st.append(["Days in data period", period_days, "Used to turn spend into spend per day"])
    st.append(["Target payment term (days)", TARGET_DAYS, "Assumption: company standard term"])
    st["B2"].number_format = st["B3"].number_format = "yyyy-mm-dd"
    for col, w in zip("ABC", [30, 12, 48]):
        st.column_dimensions[col].width = w
    header(st)
    st["A7"] = "All data is synthetic. Spend = quantity received x unit price from the purchase-order lines."
    st["A7"].font = Font(italic=True)

    wb.save(XLSX)
    print(f"{n - 1} suppliers, period {first.date()} to {last.date()} ({period_days} days) -> {XLSX}")


if __name__ == "__main__":
    main()
